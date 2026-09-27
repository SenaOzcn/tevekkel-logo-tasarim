#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Tevekkel Yedek Parça — Logo Tasarımları ve Sunumu
=========================================================
"""

import os
import sys

from svglib.svglib import svg2rlg
from reportlab.graphics import renderPDF
import pymupdf 
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


# ---------------------------------------------------------------------------
# 1) LOGO TANIMLARI
# ---------------------------------------------------------------------------

LOGOS = [
    {
        "slug": "civata",
        "name": "Cıvata",
        "tag": "Sağlam bağlantı",
        "primary": "12294D",
        "accent": "D93A2F",
        "colors": [
            ("12294D", "Gece Laciverti", "Ana renk"),
            ("D93A2F", "Sinyal Kırmızısı", "Vurgu"),
            ("FFFFFF", "Beyaz", "Zemin, ters renk"),
            ("6B7280", "Yardımcı Gri", "İkincil metin"),
        ],
        "meaning": (
            "T harfi, başı ve dişli gövdesiyle bir cıvata olarak çizildi. "
            "Cıvata sağlam bağlantıyı ve yedek parçanın kendisini simgeler. "
            "Başın altındaki kırmızı rondela hız ve dikkat vurgusu verir."
        ),
        "strength": "Sektörü hemen anlatır, tek renkte de çalışır.",
        "caution": "Hırdavat çağrışımı yapabilir; ince dişler küçük boyutta kapanır.",
        "embed": lambda p, a: f'''
            <g fill="#{p}">
                <polygon points="12,8 48,8 52,12 52,23 8,23 8,12"/>
                <polygon points="23,29 37,29 37,52 34,56 26,56 23,52"/>
            </g>
            <rect x="12" y="25" width="36" height="4" fill="#{a}"/>
            <path d="M22.7 8V23M37.3 8V23M23 38L37 34M23 44L37 40M23 50L37 46"
                  stroke="#FFFFFF" stroke-width="1.8" fill="none"/>
        ''',
    },
    {
        "slug": "piston",
        "name": "Piston",
        "tag": "Motorun gücü",
        "primary": "1F2933",
        "accent": "E8590C",
        "colors": [
            ("1F2933", "Antrasit", "Ana renk"),
            ("E8590C", "Motor Turuncusu", "Vurgu"),
            ("FFFFFF", "Beyaz", "Zemin, ters renk"),
            ("6B7280", "Yardımcı Gri", "İkincil metin"),
        ],
        "meaning": (
            "Piston başı T'nin çatısı, biyel kolu gövdesi oldu. Motoru, gücü "
            "ve mekanik ustalığı çağrıştırır. Turuncu segman bantları "
            "enerjiyi ve sıcaklığı temsil eder."
        ),
        "strength": "Otomotive en özgü sembol.",
        "caution": "Silüet boya fırçasına benzeyebilir; segman çizgileri küçük boyutta incelir.",
        "embed": lambda p, a: f'''
            <g fill="#{p}">
                <path d="M8 10Q8 6 12 6H48Q52 6 52 10V24H8Z"/>
                <polygon points="14,24 46,24 35,34 25,34"/>
                <rect x="25" y="34" width="10" height="10"/>
                <path fill-rule="evenodd"
                      d="M24 50a6 6 0 1 0 12 0a6 6 0 1 0 -12 0Z
                         M27.6 50a2.4 2.4 0 1 0 4.8 0a2.4 2.4 0 1 0 -4.8 0Z"/>
            </g>
            <g fill="#{a}">
                <rect x="8" y="9" width="44" height="4"/>
                <rect x="8" y="16" width="44" height="4"/>
            </g>
        ''',
    },
    {
        "slug": "jant",
        "name": "Jant / Lastik",
        "tag": "Yolda güvenlik",
        "primary": "111827",
        "accent": "FACC15",
        "colors": [
            ("111827", "Asfalt Siyahı", "Ana renk"),
            ("FACC15", "Uyarı Sarısı", "Vurgu, rozet zemini"),
            ("FFFFFF", "Beyaz", "Zemin"),
            ("6B7280", "Yardımcı Gri", "İkincil metin"),
        ],
        "meaning": (
            "Lastik halkasının içindeki sarı zeminde T harfi durur. Tekerlek "
            "yolu ve hareketi, sarı ise görünürlüğü ve güvenliği simgeler. "
            "Lastik ürün grubuna da doğrudan gönderme yapar."
        ),
        "strength": "Yuvarlak rozet: profil fotoğrafı, favicon ve mühür için ideal.",
        "caution": "Sarı beyaz zeminde zayıf kalır; lastik satmayan bir imaj verebilir.",
        "embed": lambda p, a: f'''
            <circle cx="30" cy="30" r="26" fill="none" stroke="#{p}" stroke-width="6"/>
            <circle cx="30" cy="30" r="22.5" fill="#{a}"/>
            <rect x="15" y="19" width="30" height="8" fill="#{p}"/>
            <rect x="26" y="27" width="8" height="18" fill="#{p}"/>
        ''',
    },
    {
        "slug": "kalkan",
        "name": "Kalkan",
        "tag": "Güvenin adresi",
        "primary": "7A1F2B",
        "accent": "C9A227",
        "colors": [
            ("7A1F2B", "Bordo", "Ana renk"),
            ("C9A227", "Altın", "Vurgu"),
            ("FFFFFF", "Beyaz", "Zemin, T harfi"),
            ("1F2937", "Koyu Gri", "Metin"),
        ],
        "meaning": (
            "Kalkan güvenceyi, beyaz T harfi markayı simgeler. Alttaki altın "
            "elmas kaliteyi ve köklülüğü (1978'den beri) temsil eder. "
            "\u201cOtomotivde Güvenin Adresi\u201d sloganıyla en uyumlu yön."
        ),
        "strength": "Slogan ve deneyim mesajını taşır, kurumsal durur.",
        "caution": "Kalkan sigorta ve güvenlik firmalarında da yaygın; sektör ipucu zayıf.",
        "embed": lambda p, a: f'''
            <path d="M30 3L54 11V30C54 44 43 52 30 58C17 52 6 44 6 30V11Z" fill="#{p}"/>
            <rect x="17" y="15" width="26" height="8" fill="#FFFFFF"/>
            <rect x="26" y="23" width="8" height="19" fill="#FFFFFF"/>
            <polygon points="30,44 35,49 30,54 25,49" fill="#{a}"/>
        ''',
    },
    {
        "slug": "hiz",
        "name": "Hız Çizgileri",
        "tag": "Hızlı tedarik",
        "primary": "0F766E",
        "accent": "F59E0B",
        "colors": [
            ("0F766E", "Petrol Yeşili", "Ana renk"),
            ("F59E0B", "Amber", "Vurgu"),
            ("FFFFFF", "Beyaz", "Zemin, ters renk"),
            ("134E4A", "Koyu Petrol", "Metin"),
        ],
        "meaning": (
            "Öne eğik T harfinin arkasında incelen üç çizgi hareketi "
            "gösterir: hızlı tedarik ve zamanında teslimat. Petrol yeşili "
            "güven ve sakinliği, amber enerjiyi anlatır."
        ),
        "strength": "Modern ve dinamik; lojistik ve B2B kimliğine uygun.",
        "caution": "Sektör bağı en zayıf yön; nakliye firmalarına da uyar.",
        "embed": lambda p, a: f'''
            <polygon points="14,10 54,10 50,22 10,22" fill="#{p}"/>
            <rect x="24" y="22" width="12" height="30" fill="#{p}"/>
            <g fill="#{a}">
                <rect x="2" y="30" width="15" height="3.5"/>
                <rect x="6" y="38" width="11" height="3.5"/>
                <rect x="10" y="46" width="7" height="3.5"/>
            </g>
        ''',
    },
]

FONT_NOTE = (
    "Kalın sans-serif (Arial/Helvetica Bold). Logo yazısı 30 pt, harf "
    "aralığı +1; alt yazı 12 pt, harf aralığı +5."
)

# Sunumda kullanılan ortak renkler
NAVY = "12294D"
RED = "D93A2F"
LIGHT_BG = "F4F6F8"
INK = "1F2937"
MUTED = "6B7280"
BORDER = "E5E7EB"
WHITE = "FFFFFF"


# ---------------------------------------------------------------------------
# 2) SVG ÜRETİMİ
# ---------------------------------------------------------------------------

def build_logo_svg(logo: dict) -> str:
    p, a = logo["primary"], logo["accent"]
    embed = logo["embed"](p, a)
    return f'''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 100"
     width="900" height="300" role="img">
  <title>Tevekkel Yedek Parça logosu — {logo["name"]}</title>
  <desc>{logo["meaning"]}</desc>
  <rect width="300" height="100" fill="#FFFFFF"/>
  <g transform="translate(10 10) scale(1.3333)">
    {embed}
  </g>
  <text x="104" y="53" font-family="Helvetica-Bold"
        font-size="30" letter-spacing="1" fill="#{p}">TEVEKKEL</text>
  <text x="105" y="76" font-family="Helvetica-Bold"
        font-size="12" letter-spacing="5" fill="#{p}">YEDEK PARÇA</text>
</svg>'''


def render_svg_to_png(svg_code: str, out_path: str, zoom: float = 4.0) -> None:
    tmp_svg = out_path + ".tmp.svg"
    tmp_pdf = out_path + ".tmp.pdf"
    with open(tmp_svg, "w", encoding="utf-8") as f:
        f.write(svg_code)
    drawing = svg2rlg(tmp_svg)
    renderPDF.drawToFile(drawing, tmp_pdf)
    doc = pymupdf.open(tmp_pdf)
    pix = doc[0].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    pix.save(out_path)
    doc.close()
    os.remove(tmp_svg)
    os.remove(tmp_pdf)


def generate_all_logos(out_dir: str) -> dict:
    paths = {}
    for i, logo in enumerate(LOGOS, start=1):
        svg_code = build_logo_svg(logo)
        svg_path = os.path.join(out_dir, f"tevekkel-logo-{i}-{logo['slug']}.svg")
        png_path = os.path.join(out_dir, f"_preview-{logo['slug']}.png")
        with open(svg_path, "w", encoding="utf-8") as f:
            f.write(svg_code)
        render_svg_to_png(svg_code, png_path)
        paths[logo["slug"]] = {"svg": svg_path, "png": png_path}
        print(f"  ✓ {svg_path}")
    return paths


# ---------------------------------------------------------------------------
# 3) PPTX YARDIMCI FONKSİYONLARI  (python-pptx)
# ---------------------------------------------------------------------------

def rgb(hex_code: str) -> RGBColor:
    return RGBColor.from_string(hex_code)


def set_background(slide, hex_code: str):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(hex_code)


def add_textbox(slide, x, y, w, h, text, size=14, bold=False, color=INK,
                 font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
                 spacing=None, line_spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = rgb(color)
    if spacing is not None:
        _set_letter_spacing(run, spacing)
    return box


def add_bullets(slide, x, y, w, h, items, size=14, color=INK, font="Calibri",
                 space_after=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        _add_bullet_char(p)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = rgb(color)
    return box


def _add_bullet_char(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    buChar = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(buChar)
    pPr.set("marL", "228600")
    pPr.set("indent", "-228600")


def _set_letter_spacing(run, points_hundredths):
    """Harf aralığı (tracking) — birim: 1/100 pt."""
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(points_hundredths))


def add_card(slide, x, y, w, h, fill=WHITE, line_color=BORDER, radius=0.08):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color)
    shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def add_image_contained(slide, path, x, y, max_w, max_h):
    """Görseli en-boy oranını koruyarak verilen kutunun içine ortalar."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    w, h = iw * ratio, ih * ratio
    left = x + (max_w - w) / 2
    top = y + (max_h - h) / 2
    slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(w), Inches(h))


def add_color_swatch_row(slide, x, y, hex_code, name, usage):
    sw = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.4), Inches(0.32))
    sw.fill.solid()
    sw.fill.fore_color.rgb = rgb(hex_code)
    sw.line.color.rgb = rgb(BORDER)
    sw.line.width = Pt(0.75)
    sw.shadow.inherit = False
    add_textbox(slide, x + 0.55, y - 0.03, 2.9, 0.38, f"{name}   #{hex_code}",
                size=13, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, x + 3.4, y - 0.03, 1.9, 0.38, usage,
                size=11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# 4) SLAYTLARIN OLUŞTURULMASI
# ---------------------------------------------------------------------------

def build_presentation(logo_paths: dict, out_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # --- Slayt 1: Kapak -----------------------------------------------
    s = prs.slides.add_slide(blank)
    set_background(s, NAVY)
    add_textbox(s, 0.8, 2.3, 6.6, 1.2, "Tevekkel Yedek Parça",
                size=40, bold=True, color=WHITE, font="Cambria")
    add_textbox(s, 0.8, 3.55, 6.6, 0.5, "Logo tasarımları · 5 farklı tasarım",
                size=20, color="CADCFC")
    add_textbox(s, 0.8, 4.1, 6.6, 0.4, "Anlamlar, renk kodları ve tipografi",
                size=14, color="CADCFC")

    # --- Slayt 2: Marka özeti ------------------------------------------
    s = prs.slides.add_slide(blank)
    set_background(s, LIGHT_BG)
    add_textbox(s, 0.6, 0.4, 12, 0.7, "Marka özeti", size=30, bold=True,
                color=NAVY, font="Cambria")
    add_bullets(s, 0.6, 1.6, 6.2, 2.8, [
        "1978'den beri Düzce merkezli otomotiv yedek parça",
        "Binek, ticari, ağır vasıta ve lastik ürün grupları",
        "Bayilere yönelik B2B sipariş sistemi",
        "Slogan: \u201cOtomotivde Güvenin Adresi\u201d",
    ], size=17)
    add_textbox(s, 0.6, 4.6, 6.2, 0.4, "TASARIM HEDEFLERİ", size=13, bold=True, color=MUTED)
    add_textbox(s, 0.6, 5.0, 6.2, 0.9,
                "Güven · hızlı tedarik · sektör netliği · küçük boyutta okunurluk",
                size=17, color=INK)
    stats = [("1978", "Kuruluş yılı"), ("4", "Ürün grubu"), ("B2B", "Bayi sistemi")]
    for k, (num, label) in enumerate(stats):
        y = 1.6 + k * 1.75
        add_card(s, 7.4, y, 5.3, 1.5)
        add_textbox(s, 7.7, y, 2.4, 1.5, num, size=40, bold=True, color=NAVY,
                    font="Cambria", anchor=MSO_ANCHOR.MIDDLE)
        add_textbox(s, 10.2, y, 2.3, 1.5, label, size=16, color=MUTED,
                    anchor=MSO_ANCHOR.MIDDLE)

    # --- Slayt 3: 5 tasarım bir arada -------------------------------------
    s = prs.slides.add_slide(blank)
    set_background(s, LIGHT_BG)
    add_textbox(s, 0.6, 0.4, 12, 0.7, "5 tasarım bir arada", size=30, bold=True,
                color=NAVY, font="Cambria")
    for i, logo in enumerate(LOGOS):
        x = 0.6 + (i % 3) * 4.15
        y = 1.4 + (i // 3) * 2.75
        add_card(s, x, y, 3.9, 2.5)
        add_image_contained(s, logo_paths[logo["slug"]]["png"], x + 0.15, y + 0.2, 3.6, 1.2)
        add_textbox(s, x + 0.25, y + 1.55, 3.4, 0.35, f"{i+1}. {logo['name']}",
                    size=15, bold=True, color=NAVY)
        add_textbox(s, x + 0.25, y + 1.9, 3.4, 0.3, logo["tag"], size=12, color=MUTED)
    add_card(s, 8.9, 4.15, 3.9, 2.5, fill=NAVY, line_color=NAVY)
    add_textbox(s, 9.15, 4.4, 3.4, 2.0,
                "Karşılaştırmayı kolaylaştırmak için beş tasarımda da yazı düzeni ve "
                "font aynı tutuldu; fark yalnızca sembolden gelir.",
                size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    # --- Slayt 4-8: Her logo için detay sayfası --------------------------
    for i, logo in enumerate(LOGOS):
        s = prs.slides.add_slide(blank)
        set_background(s, LIGHT_BG)
        add_textbox(s, 0.6, 0.35, 8, 0.6, f"{i+1}. {logo['name']}", size=28,
                    bold=True, color=NAVY, font="Cambria")
        add_textbox(s, 0.6, 0.95, 8, 0.35, logo["tag"], size=15, color=MUTED)

        add_card(s, 0.6, 1.5, 6.6, 3.2)
        add_image_contained(s, logo_paths[logo["slug"]]["png"], 0.85, 1.75, 6.1, 2.7)

        # Güçlü yön / Dikkat kartları
        add_card(s, 0.6, 5.0, 3.2, 1.9)
        add_textbox(s, 0.8, 5.12, 2.8, 0.3, "Güçlü yön", size=13, bold=True, color=NAVY)
        add_textbox(s, 0.8, 5.45, 2.8, 1.3, logo["strength"], size=13, color=INK)

        add_card(s, 4.0, 5.0, 3.2, 1.9)
        add_textbox(s, 4.2, 5.12, 2.8, 0.3, "Dikkat", size=13, bold=True, color=RED)
        add_textbox(s, 4.2, 5.45, 2.8, 1.3, logo["caution"], size=13, color=INK)

        # Sağ sütun: anlam, renk paleti, tipografi
        add_textbox(s, 7.6, 1.5, 5.1, 0.35, "ANLAM", size=13, bold=True, color=MUTED)
        add_textbox(s, 7.6, 1.9, 5.1, 1.5, logo["meaning"], size=14, color=INK)

        add_textbox(s, 7.6, 3.5, 5.1, 0.35, "RENK PALETİ", size=13, bold=True, color=MUTED)
        for k, (hex_code, cname, usage) in enumerate(logo["colors"]):
            add_color_swatch_row(s, 7.6, 3.9 + k * 0.5, hex_code, cname, usage)

        add_textbox(s, 7.6, 5.95, 5.1, 0.35, "TİPOGRAFİ", size=13, bold=True, color=MUTED)
        add_textbox(s, 7.6, 6.3, 5.1, 0.6, FONT_NOTE, size=13, color=INK)

    # --- Değerlendirme ve Karar ------------------------------
    s = prs.slides.add_slide(blank)
    set_background(s, NAVY)
 
    add_textbox(s, 1.0, 2.05, 11.33, 0.9, "Değerlendirme ve Karar",
                size=38, bold=True, color=WHITE, font="Calibri",
                align=PP_ALIGN.CENTER)
    add_textbox(s, 2.0, 2.95, 9.33, 0.9,
                "Beş logo tasarımından markanıza en uygun olanı seçtikten sonra "
                "renk, kurumsal evrak ve dijital uyarlamalar tamamlanacaktır.",
                size=15, color="9CA6BB", font="Calibri",
                align=PP_ALIGN.CENTER, line_spacing=1.25)
 
    card_w, card_h, gap = 3.35, 1.55, 0.3
    total_w = card_w * 2 + gap
    x1 = (13.333 - total_w) / 2
    x2 = x1 + card_w + gap
    y = 4.35
    card_fill, card_line = "1E3A66", "3B5A8F"
 
    add_card(s, x1, y, card_w, card_h, fill=card_fill, line_color=card_line, radius=0.14)
    add_textbox(s, x1 + 0.25, y + 0.28, card_w - 0.5, 0.3, "SONRAKİ ADIM",
                size=11, bold=True, color="8FA8D9", font="Calibri",
                align=PP_ALIGN.CENTER, spacing=100)
    add_textbox(s, x1 + 0.25, y + 0.68, card_w - 0.5, 0.6, "Logo Tasarımının Seçilmesi",
                size=16, bold=True, color=WHITE, font="Calibri",
                align=PP_ALIGN.CENTER)
 
    add_card(s, x2, y, card_w, card_h, fill=card_fill, line_color=card_line, radius=0.14)
    add_textbox(s, x2 + 0.25, y + 0.28, card_w - 0.5, 0.3, "TESLİMAT",
                size=11, bold=True, color="8FA8D9", font="Calibri",
                align=PP_ALIGN.CENTER, spacing=100)
    add_textbox(s, x2 + 0.25, y + 0.68, card_w - 0.5, 0.6, "Vektörel Formatlar (SVG, PDF)",
                size=16, bold=True, color=WHITE, font="Calibri",
                align=PP_ALIGN.CENTER)

    prs.save(out_path)
    ok = os.path.exists(out_path)
    size = os.path.getsize(out_path) if ok else 0
    print(f"  {'✓' if ok else '✗'} {os.path.abspath(out_path)}  (var mı: {ok}, boyut: {size} bayt)")


# ---------------------------------------------------------------------------
# 5) ANA AKIŞ
# ---------------------------------------------------------------------------

def main():
    out_dir = sys.argv[1] if len(sys.argv) > 1 else "tevekkel_output"
    os.makedirs(out_dir, exist_ok=True)

    print("Logolar (SVG) üretiliyor...")
    logo_paths = generate_all_logos(out_dir)

    print("Sunum (PPTX) oluşturuluyor...")
    pptx_path = os.path.join(out_dir, "Tevekkel-Logo-Sunumu.pptx")
    try:
        build_presentation(logo_paths, pptx_path)
    except Exception:
        import traceback
        print("\n!!! Sunum oluşturulurken bir hata oluştu !!!")
        traceback.print_exc()
        print("\nYukarıdaki hata metnini olduğu gibi paylaşın.")
        sys.exit(1)

    for slug, p in logo_paths.items():
        if os.path.exists(p["png"]):
            os.remove(p["png"])

    print("\nTamamlandı. Çıktılar:", os.path.abspath(out_dir))


if __name__ == "__main__":
    main()
